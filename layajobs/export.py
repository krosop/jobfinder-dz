"""Excel + JSONL exports of stored offers."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_COLUMNS = [
    ("id", 6), ("source_newspaper", 20), ("job_title", 40),
    ("organization_or_company", 40), ("sector", 10), ("contract_type", 13),
    ("wilaya_location", 16), ("positions_count", 9),
    ("required_qualifications", 50), ("application_deadline", 14),
    ("application_deadline_raw", 26), ("deadline_kind", 11),
    ("contact_email", 26), ("contact_phone", 16),
    ("contact_address_or_portal", 36), ("language", 9),
    ("times_seen", 10), ("first_seen", 20), ("last_seen", 20),
    ("gate", 8), ("page", 7), ("source_file", 30), ("snippet", 60),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_row(ws, row_idx: int, values: List[Any], *, header: bool = False) -> None:
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if header:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
        else:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_excel(rows: List[Dict[str, Any]], out_path: str | Path) -> Path:
    """Write offers into an xlsx: a Summary sheet + one sheet per source."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Export", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    ws.append(["Total offers", len(rows)])
    ws.append([])
    by_source: Dict[str, int] = {}
    for r in rows:
        key = r.get("source_newspaper") or r.get("source_slug") or "?"
        by_source[key] = by_source.get(key, 0) + 1
    ws.append(["Source", "Offers"])
    _write_row(ws, ws.max_row, ["Source", "Offers"], header=True)
    for k in sorted(by_source, key=lambda x: -by_source[x]):
        ws.append([k, by_source[k]])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 10

    # One sheet per source
    by_slug: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        slug = r.get("source_slug") or "other"
        by_slug.setdefault(slug, []).append(r)

    headers = [c for c, _ in _COLUMNS]
    for slug, group in sorted(by_slug.items()):
        safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in slug)[:28]
        ws = wb.create_sheet(title=safe)
        _write_row(ws, 1, headers, header=True)
        for i, r in enumerate(group, start=2):
            values = []
            for col, _w in _COLUMNS:
                v = r.get(col)
                if col == "required_qualifications" and isinstance(v, str):
                    try:
                        v = "\n".join(json.loads(v)) if v else ""
                    except (json.JSONDecodeError, ValueError):
                        pass
                values.append(v)
            _write_row(ws, i, values)
        for j, (_col, width) in enumerate(_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.freeze_panes = "A2"

    wb.save(out_path)
    return out_path


def export_jsonl(rows: List[Dict[str, Any]], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False, default=str) + "\n")
    return out_path
