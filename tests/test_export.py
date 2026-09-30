"""Tests for layajobs.export — Excel + JSONL outputs."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from layajobs.export import export_excel, export_jsonl
from layajobs.store import Store


@pytest.fixture()
def rows(tmp_path):
    st = Store(tmp_path / "db.sqlite")
    offers = [
        {"source_newspaper": "El Test", "source_slug": "el-test",
         "job_title": "Ingénieur", "organization_or_company": "SARL X",
         "sector": "Private", "contract_type": "CDI",
         "wilaya_location": "Blida", "positions_count": 2,
         "required_qualifications": ["Diplôme d'ingénieur", "3 ans d'expérience"],
         "application_deadline": "2026-10-25", "application_deadline_raw": "avant le 25 octobre 2026",
         "deadline_kind": "absolute",
         "contact_info": {"email": "a@b.dz", "phone": "025410892",
                          "address_or_portal": "ZI Ouled Yaich"},
         "language": "fr", "snippet": "AVIS DE RECRUTEMENT ...",
         "laya_confidence": {"sector": 0.93}, "gate": "laya"},
        {"source_newspaper": "جريدة اختبار", "source_slug": "ar-test",
         "job_title": "مهندس دولة", "organization_or_company": "مديرية الأشغال",
         "sector": "Public", "contract_type": "Statutaire",
         "wilaya_location": "Batna", "positions_count": None,
         "required_qualifications": [],
         "application_deadline": None, "application_deadline_raw": "15 يوما",
         "deadline_kind": "relative",
         "contact_info": {"email": None, "phone": None, "address_or_portal": None},
         "language": "ar", "snippet": "إعلان توظيف ...",
         "laya_confidence": {}, "gate": "keyword_fallback"},
    ]
    for o in offers:
        st.upsert_offer(o)
    return st.query()


def test_excel_structure(rows, tmp_path):
    out = export_excel(rows, tmp_path / "reports" / "t.xlsx")
    assert out.exists()
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert "el-test" in wb.sheetnames and "ar-test" in wb.sheetnames
    ws = wb["el-test"]
    header = [c.value for c in ws[1]]
    assert "job_title" in header
    assert ws.max_row == 2  # header + one offer
    # qualifications JSON string rendered as multi-line text
    idx = header.index("required_qualifications") + 1
    assert "Diplôme d'ingénieur" in ws.cell(row=2, column=idx).value


def test_jsonl_roundtrip(rows, tmp_path):
    out = export_jsonl(rows, tmp_path / "t.jsonl")
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == len(rows)
    rec = json.loads(lines[0])
    assert rec["job_title"] in ("Ingénieur", "مهندس دولة")
    assert "laya_confidence" in rec
